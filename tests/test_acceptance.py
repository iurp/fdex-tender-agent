"""Checks the public corpus and harness, not unimplemented learner features.

Oracle-derived answers here test the checker itself. No product runtime imports
them, and these tests must never be reported as student module completion.
"""
import copy
import importlib.util
import json
import contextlib
import io
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("task_acceptance", ROOT / "tools/task_acceptance.py")
harness = importlib.util.module_from_spec(spec); spec.loader.exec_module(harness)
sys.path.insert(0, str(ROOT / "tools"))
from run_pr_acceptance import assigned_module


def write(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


class AcceptanceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="tender-kit-test-")
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def manifest(self, module, case="normal"):
        return next(c for c in harness.cases(module) if c["case_id"] == case)

    def oracle_projection(self, manifest):
        oracle = harness.read(ROOT / manifest["expected"])
        return {"schema_version": harness.RESULT_VERSION, "module": manifest["module"], "case_id": manifest["case_id"], "source_mode": "DEMO", **copy.deepcopy(oracle),
                "input_hashes": {p["path"]: harness.digest(ROOT / p["path"]) for p in manifest["inputs"]}, "evidence": [], "artifacts": {}}

    def add_test_artifacts(self, manifest, result, directory):
        """Construct checker-positive golden outputs, NOT production solutions."""
        request = harness.read(ROOT / manifest["input"])["payload"]
        if manifest["module"] == "M04" and manifest["case_id"] == "normal":
            result["evidence"] = [{"path": request["document"], "kind": "line", "line": r["line"], "quote": r["quote"]} for r in result["facts"]["requirements"]]
        if manifest["module"] == "M08" and manifest["case_id"] == "normal":
            for a in request["assets"]:
                for f, box in (("person_id", [65, 175, 1250, 245]), ("certificate_id", [65, 280, 1250, 350]), ("qualification", [65, 385, 1250, 455]), ("valid_until", [65, 490, 1250, 560])):
                    result["evidence"].append({"path": a["image"], "kind": "image_region", "field": f, "box": box})
        artifact = manifest["artifacts"]
        kind = artifact.get("kind")
        if kind == "personnel_docx":
            doc = Document(ROOT / request["template"]["path"])
            for p in doc.paragraphs:
                for run in p.runs:
                    run.text = run.text.replace("{{PROJECT_NAME}}", "示例专业服务项目").replace("{{PROJECT_ID}}", "DEMO-BID-001").replace("{{LOT_ID}}", "LOT-001")
            for values in artifact["rows"]:
                row = doc.tables[0].add_row()
                for cell, value in zip(row.cells, values):
                    cell.text = value
            for image in artifact["image_paths"]:
                doc.add_page_break(); doc.add_picture(str(ROOT / image), width=Inches(6.1))
            doc.save(directory / "actual.docx")
            result["artifacts"]["document"] = "actual.docx"
        if kind == "public_zip":
            with zipfile.ZipFile(directory / "public.zip", "w") as z:
                for path in artifact["allowed_inputs"]:
                    z.write(ROOT / path, Path(path).name)
            result["artifacts"]["archive"] = "public.zip"
        if kind in ("rerun_history", "rerun_scenarios"):
            scenarios = result["facts"].get("scenarios", {"normal": result["facts"]})
            for name, facts in scenarios.items():
                old, new = directory / (name + "-old"), directory / (name + "-new")
                old.mkdir(); new.mkdir()
                for module, item in request["previous_outputs"].items():
                    shutil.copy2(ROOT / item["path"], old / f"{module}.json")
                    if module in facts["unchanged"]:
                        shutil.copy2(ROOT / item["path"], new / f"{module}.json")
                    else:
                        write(new / f"{module}.json", {"module": module, "run_id": request["new_run_id"], "status": "COMPLETE"})
                write(new / "release_record.json", {"status": "STALE"})
                refs = {"old_run": old.name, "new_run": new.name}
                if kind == "rerun_history":
                    result["artifacts"] = refs
                else:
                    result["artifacts"].setdefault("scenarios", {})[name] = refs

    def check(self, manifest, result, directory=None):
        directory = directory or self.root
        return harness.check_result(manifest, result, directory / "result.json")

    def test_all_ready_tasks_have_four_categories_and_hashes(self):
        errors, count = harness.verify_kit()
        self.assertEqual(errors, [])
        self.assertEqual(count, 54)
        ready = {t["module"] for t in harness.read(ROOT / "configs/tasks.json") if t["status"] == "ready"}
        self.assertEqual(ready, {c["module"] for c in harness.cases()})

    def test_all_54_oracles_are_checker_positive_not_learner_passes(self):
        for manifest in harness.cases():
            with self.subTest(module=manifest["module"], case=manifest["case_id"]):
                directory = self.root / (manifest["module"] + "-" + manifest["case_id"]); directory.mkdir()
                result = self.oracle_projection(manifest)
                self.add_test_artifacts(manifest, result, directory)
                self.assertEqual(self.check(manifest, result, directory), [])

    def test_every_case_rejects_wrong_status_wrong_hash_and_real_authority(self):
        for manifest in harness.cases():
            with self.subTest(module=manifest["module"], case=manifest["case_id"]):
                result = self.oracle_projection(manifest)
                result["status"] = "COMPLETE" if result["status"] != "COMPLETE" else "BLOCKED"
                self.assertTrue(self.check(manifest, result))
                result = self.oracle_projection(manifest); result["input_hashes"][manifest["input"]] = "0" * 64
                self.assertTrue(any("input_hashes" in e for e in self.check(manifest, result)))
                result = self.oracle_projection(manifest); result["real_submission_authorized"] = True
                self.assertTrue(any("real_submission_authorized" in e for e in self.check(manifest, result)))

    def test_each_oracle_fact_cannot_be_missing(self):
        for manifest in harness.cases():
            for key in harness.read(ROOT / manifest["expected"])["facts"]:
                with self.subTest(module=manifest["module"], case=manifest["case_id"], key=key):
                    result = self.oracle_projection(manifest); del result["facts"][key]
                    self.assertTrue(any("facts." + key in e for e in self.check(manifest, result)))

    def test_number_strings_booleans_null_and_order_are_distinct(self):
        self.assertTrue(harness.compare(False, 0))
        self.assertTrue(harness.compare("318.00", 318.0))
        self.assertTrue(harness.compare(None, "unknown"))
        self.assertTrue(harness.compare(["M09", "M13"], ["M13", "M09"]))

    def test_pricing_oracle_is_independently_recomputed_from_input(self):
        manifest = self.manifest("M12")
        lines = harness.read(ROOT / manifest["input"])["payload"]["lines"]
        expected = harness.read(ROOT / manifest["expected"])["facts"]["quotes"]
        for line in lines:
            quantize = lambda x: x.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            amount = quantize(Decimal(line["quantity"]) * Decimal(line["unit_price"]))
            rate = Decimal(line["tax_rate"])
            if line["tax_mode"] == "exclusive":
                net = amount; tax = quantize(net * rate); gross = net + tax
            else:
                gross = amount; net = quantize(gross / (Decimal("1") + rate)); tax = gross - net
            oracle = expected[line["id"]]
            self.assertEqual((str(net), str(tax), str(gross)), (oracle["net"], oracle["tax"], oracle["gross"]))

    def test_missing_issue_codes_fail_all_review_and_abnormal_cases(self):
        for manifest in harness.cases():
            result = self.oracle_projection(manifest)
            if result["issue_codes"]:
                result["issue_codes"] = []
                self.assertTrue(any("issue_codes" in e for e in self.check(manifest, result)))

    def test_plain_text_has_no_teacher_tags_and_exact_12_lines(self):
        manifest = self.manifest("M04")
        request = harness.read(ROOT / manifest["input"])
        body = (ROOT / request["payload"]["document"]).read_text()
        self.assertNotIn("[R-", body)
        result = self.oracle_projection(manifest); self.add_test_artifacts(manifest, result, self.root)
        result["evidence"][0]["line"] = 1
        self.assertTrue(any("evidence" in e.lower() for e in self.check(manifest, result)))

    def test_actual_ocr_images_readable_and_date_redacted(self):
        normal = harness.read(ROOT / self.manifest("M08")["input"])["payload"]
        for asset in normal["assets"]:
            with Image.open(ROOT / asset["image"]) as image:
                self.assertEqual(image.size, (1400, 850))
                image.verify()
        review = harness.read(ROOT / self.manifest("M08", "review")["input"])["payload"]
        with Image.open(ROOT / review["assets"][0]["image"]) as image:
            self.assertEqual(image.getpixel((1300, 590)), (176, 176, 176))
        result = self.oracle_projection(self.manifest("M08")); self.add_test_artifacts(self.manifest("M08"), result, self.root)
        result["evidence"][0]["box"] = [0, 0, 9999, 9000]
        self.assertTrue(any("box" in e for e in self.check(self.manifest("M08"), result)))

    def test_malformed_evidence_does_not_crash_or_pass(self):
        for module in ("M04", "M08"):
            manifest = self.manifest(module); result = self.oracle_projection(manifest)
            result["evidence"] = ["not a record"]
            self.assertTrue(self.check(manifest, result))

    def test_word_media_set_cannot_hide_swapped_relationships(self):
        original = ROOT / "samples/acceptance/shared/word/personnel-chapter.docx"
        swapped = ROOT / "samples/acceptance/shared/word/swap-images.docx"
        with zipfile.ZipFile(original) as a, zipfile.ZipFile(swapped) as b:
            self.assertEqual({harness.hashlib.sha256(a.read(n)).hexdigest() for n in a.namelist() if n.startswith("word/media/")},
                             {harness.hashlib.sha256(b.read(n)).hexdigest() for n in b.namelist() if n.startswith("word/media/")})
        self.assertNotEqual(harness.image_sequence(Document(original)), harness.image_sequence(Document(swapped)))

    def test_docx_claim_without_actual_file_fails(self):
        manifest = self.manifest("M13"); result = self.oracle_projection(manifest)
        result["artifacts"] = {"document": "missing.docx"}
        self.assertTrue(self.check(manifest, result))

    def test_actual_word_wrong_cell_and_old_person_fail(self):
        manifest = self.manifest("M13"); result = self.oracle_projection(manifest)
        self.add_test_artifacts(manifest, result, self.root)
        doc = Document(self.root / "actual.docx"); doc.tables[0].cell(1, 5).text = "错误资格"
        doc.save(self.root / "actual.docx")
        self.assertTrue(any("table" in e for e in self.check(manifest, result)))
        doc.tables[0].cell(1, 5).text = "示例专业服务"; doc.add_paragraph("旧姓名 示例人员丁")
        doc.save(self.root / "actual.docx")
        self.assertTrue(any("Old/unselected" in e for e in self.check(manifest, result)))

    def test_nonpublic_and_unconfirmed_cannot_emit_artifact(self):
        manifest = self.manifest("M18", "abnormal"); result = self.oracle_projection(manifest)
        with zipfile.ZipFile(self.root / "leaked.zip", "w") as z:
            z.writestr("fake.txt", "not real data, but should still be refused")
        self.assertTrue(any("public ZIP" in e for e in self.check(manifest, result)))
        manifest = self.manifest("M13", "review"); result = self.oracle_projection(manifest)
        shutil.copy2(ROOT / "templates/personnel-v1.docx", self.root / "unconfirmed.docx")
        self.assertTrue(any("produced DOCX" in e for e in self.check(manifest, result)))

    def prohibited_artifact_cases(self):
        return (("M13", "review", "document", ".docx", "produced DOCX"),
                ("M18", "abnormal", "archive", ".zip", "public ZIP"))

    def put_prohibited_artifact(self, path, suffix):
        path.parent.mkdir(parents=True, exist_ok=True)
        if suffix == ".docx":
            shutil.copy2(ROOT / "templates/personnel-v1.docx", path)
        else:
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("synthetic.txt", "Public synthetic regression fixture")

    def test_prohibited_artifact_suffixes_are_case_insensitive(self):
        for module, case, key, suffix, message in self.prohibited_artifact_cases():
            for index, extension in enumerate((suffix, suffix.upper(), suffix.title())):
                with self.subTest(module=module, extension=extension):
                    directory = self.root / f"{module}-{index}"; directory.mkdir()
                    manifest = self.manifest(module, case)
                    result = self.oracle_projection(manifest)
                    self.put_prohibited_artifact(directory / "nested" / ("output" + extension), suffix)
                    self.assertTrue(any(message in e for e in harness.check_artifacts(manifest, result, directory / "result.json")))

    def test_prohibited_artifact_declarations_are_checked_without_file(self):
        for module, case, key, suffix, message in self.prohibited_artifact_cases():
            for artifacts in ({key: "output.bin"}, {"extra": {"files": ["missing" + suffix.upper()]}},
                              {"extra": {key: "missing.bin"}}):
                with self.subTest(module=module, artifacts=artifacts):
                    manifest = self.manifest(module, case)
                    result = self.oracle_projection(manifest); result["artifacts"] = artifacts
                    self.assertTrue(any(message in e for e in harness.check_artifacts(manifest, result, self.root / "result.json")))

    def test_prohibited_artifact_declarations_reject_invalid_paths(self):
        for module, case, key, suffix, message in self.prohibited_artifact_cases():
            manifest = self.manifest(module, case)
            for path in ("../outside.json", str(self.root / "outside.json"), None, 12):
                with self.subTest(module=module, path=path):
                    result = self.oracle_projection(manifest); result["artifacts"] = {"summary": path}
                    self.assertTrue(self.check(manifest, result))
            result = self.oracle_projection(manifest); result["artifacts"] = {"summary": "summary.json"}
            write(self.root / "summary.json", {"status": "BLOCKED"})
            self.assertEqual(self.check(manifest, result), [])

    def test_check_cli_rejects_prohibited_files_and_declarations(self):
        for module, case, key, suffix, message in self.prohibited_artifact_cases():
            for index, variant in enumerate((suffix, suffix.upper(), suffix.title(), "declared-role", "declared-suffix")):
                with self.subTest(module=module, variant=variant):
                    directory = self.root / f"{module}-{index}"; directory.mkdir()
                    manifest = self.manifest(module, case); result = self.oracle_projection(manifest)
                    if variant == "declared-role":
                        result["artifacts"] = {key: "missing.bin"}
                    elif variant == "declared-suffix":
                        result["artifacts"] = {"extra": {"files": ["missing" + suffix.upper()]}}
                    else:
                        self.put_prohibited_artifact(directory / "nested" / ("output" + variant), suffix)
                    write(directory / "result.json", result)
                    process = subprocess.run([sys.executable, str(ROOT / "tools/task_acceptance.py"), "check",
                                              "--module", module, "--case", case, "--result", str(directory / "result.json")],
                                             cwd=ROOT, capture_output=True, text=True)
                    self.assertEqual(process.returncode, 1, process.stdout + process.stderr)
                    value = json.loads(process.stdout)
                    self.assertFalse(value["passed"])
                    self.assertTrue(any(message in e for e in value["errors"]))

    def test_zip_extra_file_and_changed_bytes_fail(self):
        manifest = self.manifest("M18"); result = self.oracle_projection(manifest)
        self.add_test_artifacts(manifest, result, self.root)
        with zipfile.ZipFile(self.root / "public.zip", "a") as z:
            z.writestr("unauthorized.txt", "synthetic forbidden extra")
        self.assertTrue(self.check(manifest, result))

    def test_history_overwrite_and_inherited_release_fail(self):
        manifest = self.manifest("X01"); result = self.oracle_projection(manifest)
        self.add_test_artifacts(manifest, result, self.root)
        old = self.root / result["artifacts"]["old_run"]
        new = self.root / result["artifacts"]["new_run"]
        write(old / "M03.json", {"overwritten": True})
        write(new / "release_record.json", {"status": "RELEASED"})
        errors = self.check(manifest, result)
        self.assertTrue(any("overwritten" in e for e in errors))
        self.assertTrue(any("inherited" in e for e in errors))

    def test_paths_cannot_escape_or_follow_external_symlink(self):
        for path in ("../outside.docx", str(ROOT / "templates/personnel-v1.docx")):
            with self.assertRaises(ValueError):
                harness.contained(self.root, path)
        (self.root / "outside").symlink_to(ROOT / "templates/personnel-v1.docx")
        with self.assertRaises(ValueError):
            harness.contained(self.root, "outside")

    def test_adapter_missing_and_output_overwrite_fail(self):
        with self.assertRaises(ValueError):
            harness.run_suite("M12", self.root / "missing.py", self.root / "out")
        adapter = self.root / "empty.py"; adapter.write_text("raise SystemExit(0)\n")
        with self.assertRaises(ValueError):
            harness.run_suite("M12", adapter, self.root)

    def test_runner_exit_zero_without_result_is_not_success(self):
        adapter = self.root / "empty.py"; adapter.write_text("raise SystemExit(0)\n")
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertFalse(harness.run_suite("M12", adapter, self.root / "out"))
        report = harness.read(self.root / "out/report.json")
        self.assertEqual(len(report["cases"]), 4)
        self.assertFalse(report["passed"])
        self.assertEqual(report["business_uat"], "PENDING")

    def test_verify_cli_explicitly_reports_student_not_tested(self):
        process = subprocess.run([sys.executable, str(ROOT / "tools/task_acceptance.py"), "verify-kit"], cwd=ROOT, capture_output=True, text=True)
        self.assertEqual(process.returncode, 0, process.stderr)
        value = json.loads(process.stdout)
        self.assertEqual(value["student_implementation"], "NOT_TESTED")
        self.assertEqual(value["cases"], 54)

    def test_ci_requires_all_eleven_task_adapters(self):
        mapping = harness.read(ROOT / "configs/task-issues.json")
        for number, module in mapping.items():
            self.assertEqual(assigned_module(f"Implement scoped feature\n\nCloses #{number}\n", "learner"), module)
        self.assertEqual(set(mapping.values()), set(harness.read(harness.KIT / "index.json")["modules"]))
        self.assertIsNone(assigned_module("Closes #22", "iurp"))
        for body, author in (("Closes #22", "learner"), ("Closes #1", "iurp"), ("Closes #3\nCloses #4", "learner"), ("no task", "learner")):
            with self.assertRaises(ValueError):
                assigned_module(body, author)


if __name__ == "__main__":
    unittest.main()
