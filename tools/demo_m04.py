"""Generate synthetic raw inputs and real M03 -> M04 handoffs for M04-01."""
import argparse
import json
import shutil
from pathlib import Path

from docx import Document
from fdex_tender.core import digest_file, read_json, save_module, write_json
from fdex_tender.modules import read_tender
from fdex_tender.pipeline import run

ROOT = Path(__file__).resolve().parents[1]


def demonstrate(output):
    output = Path(output).resolve()
    if output.exists():
        raise ValueError("Output already exists; use a new directory")
    base = ROOT / "samples/e2e/core-normal"
    plain = (ROOT / "samples/modules/M04/plain.md").read_text(encoding="utf-8")
    normal_input = output / "inputs/normal"
    shutil.copytree(base, normal_input)
    (normal_input / "tender/tender.md").write_text(plain, encoding="utf-8")
    normal_output = output / "normal"
    run(normal_input, normal_output)
    # Read the actually saved production output for the next version's history.
    previous = read_json(normal_output / "tender_brief.json")
    statuses = {"normal": previous["processing_status"]}
    for name, text in (("review", plain + "\n补充要求：人员人数恰为2名，上述人数要求仍保留。\n"),
                       ("abnormal", "# 教学文件节选\n[正文文件读取失败，未提供任何可验证的要求]\n"),
                       ("change", plain.replace("人员不少于3名", "人员不少于4名"))):
        source = output / "inputs" / name
        (source / "tender").mkdir(parents=True)
        (source / "tender/tender.md").write_text(text, encoding="utf-8")
        write_json(source / "context.json", read_json(base / "context.json"))
        manifest = read_json(base / "sources.json")
        manifest["sources"] = [s for s in manifest["sources"] if s["kind"] == "tender_markdown"]
        if name == "change":
            manifest["sources"][0]["version"] = "1.1"
        write_json(source / "sources.json", manifest)
        upstream = output / "upstream" / name
        run(source, upstream, stop="M03")
        registered_path = upstream / "source_manifest.json"
        registered = read_json(registered_path)
        ctx = read_json(upstream / ".context.json")
        value = read_tender(ctx, registered, previous if name == "change" else None)
        save_module(output / name, "tender_brief.json", value)
        write_json(output / name / "handoff.json", {
            "actual_upstream": str(registered_path), "sha256": digest_file(registered_path),
            "previous_brief": str(normal_output / "tender_brief.json") if name == "change" else None,
            "scope": "M03 registration -> M04 requirements; downstream not executed for this scenario"})
        statuses[name] = value["processing_status"]
    document = Document(normal_output / "personnel-chapter.docx")
    report = {"source_mode": "DEMO", "statuses": statuses,
              "normal_full_chain": read_json(normal_output / "run_record.json")["modules"],
              "actual_docx": {"tables": len(document.tables), "rows": [len(t.rows) for t in document.tables],
                              "inline_images": len(document.inline_shapes), "visual_review": "NOT_PERFORMED"},
              "workbuddy_runtime": "NOT_TESTED", "real_submission_authorized": False}
    write_json(output / "demo-report.json", report)
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    print(json.dumps(demonstrate(args.output), ensure_ascii=False, indent=2))
