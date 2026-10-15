#!/usr/bin/env python3
"""Offline task corpus integrity, output assertions and opt-in adapter runner.

This is a test harness, NOT an implementation of the assigned modules.
It alone reads expected.json. A learner adapter only receives input.json.
"""
import argparse
import hashlib
import json
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.opc.exceptions import PackageNotFoundError

ROOT = Path(__file__).resolve().parents[1]
KIT = ROOT / "samples/acceptance"
RESULT_VERSION = "bid-training.acceptance-result.v1"


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def contained(root, relative):
    if not isinstance(relative, str) or not relative or Path(relative).is_absolute():
        raise ValueError(f"Expected a nonempty relative path: {relative!r}")
    target = (Path(root) / relative).resolve()
    if not target.is_relative_to(Path(root).resolve()):
        raise ValueError(f"Path escapes declared root: {relative}")
    return target


def cases(module=None):
    index = read(KIT / "index.json")
    if module and module not in index["modules"]:
        raise ValueError(f"Unknown or not-open module: {module}")
    return [read(contained(ROOT, c["manifest"])) for c in index["cases"] if module is None or c["module"] == module]


def verify_case(manifest):
    failures = []
    paths = [i["path"] for i in manifest["inputs"]]
    if len(paths) != len(set(paths)) or manifest["input"] not in paths:
        failures.append("Input inventory has duplicates or lacks input.json")
    for item in manifest["inputs"]:
        try:
            file = contained(ROOT, item["path"])
            if not file.is_file() or digest(file) != item["sha256"]:
                failures.append("Input byte hash mismatch: " + item["path"])
        except (ValueError, OSError) as exc:
            failures.append(str(exc))
    request = read(contained(ROOT, manifest["input"]))
    expected = read(contained(ROOT, manifest["expected"]))
    if (request["module"], request["case_id"]) != (manifest["module"], manifest["case_id"]) or request["source_mode"] != "DEMO":
        failures.append("Request context does not match case")
    if not isinstance(expected["facts"], dict) or expected["status"] not in ("COMPLETE", "REVIEW", "BLOCKED"):
        failures.append("Malformed oracle")
    if expected["real_submission_authorized"] is not False:
        failures.append("Training corpus cannot authorize real submission")
    return failures


def verify_kit(module=None):
    failures = []
    sums = read(KIT / "SHA256SUMS.json")
    actual = {p.relative_to(ROOT).as_posix() for p in KIT.rglob("*") if p.is_file() and p.name != "SHA256SUMS.json"}
    if actual != set(sums):
        failures.append("Corpus inventory mismatch (added or missing file)")
    for name, value in sums.items():
        path = contained(ROOT, name)
        if not path.is_file() or digest(path) != value:
            failures.append("Corpus byte hash mismatch: " + name)
    selected = cases(module)
    for manifest in selected:
        failures += [f"{manifest['module']}/{manifest['case_id']}: {f}" for f in verify_case(manifest)]
    modules = {m["module"] for m in selected}
    for m in modules:
        categories = {c["category"] for c in selected if c["module"] == m}
        if not {"normal", "review", "abnormal", "change"} <= categories:
            failures.append(f"{m}: lacks a required case category")
    ready = {t["module"] for t in read(ROOT / "configs/tasks.json") if t["status"] == "ready"}
    if ready != set(read(KIT / "index.json")["modules"]):
        failures.append("Ready task coverage differs from task configuration")
    return failures, len(selected)


def compare(expected, actual, location="facts"):
    """Expected mappings are subsets; arrays and scalar types are strict."""
    errors = []
    if isinstance(expected, dict):
        if not isinstance(actual, dict):
            return [f"{location}: expected object, received {type(actual).__name__}"]
        for key, value in expected.items():
            if key not in actual:
                errors.append(f"{location}.{key}: missing")
            else:
                errors += compare(value, actual[key], f"{location}.{key}")
    elif isinstance(expected, list):
        if not isinstance(actual, list) or len(actual) != len(expected):
            return [f"{location}: array shape mismatch, expected {expected!r}"]
        for i, (a, b) in enumerate(zip(expected, actual)):
            errors += compare(a, b, f"{location}[{i}]")
    elif type(expected) is not type(actual) or expected != actual:
        errors.append(f"{location}: expected {expected!r}, received {actual!r}")
    return errors


def check_evidence(manifest, result):
    errors = []
    evidence = result.get("evidence", [])
    if not isinstance(evidence, list):
        return ["evidence must be an array"]
    allowed = {i["path"] for i in manifest["inputs"]}
    for item in evidence:
        if not isinstance(item, dict) or item.get("path") not in allowed:
            errors.append("Evidence refers to an unregistered file")
            continue
        kind = item.get("kind")
        if kind == "line":
            try:
                lines = contained(ROOT, item["path"]).read_text(encoding="utf-8").splitlines()
                number = item["line"]
                if type(number) is not int or number < 1 or number > len(lines) or item.get("quote") != lines[number - 1]:
                    errors.append("Line evidence does not match actual source text")
            except (OSError, UnicodeError, KeyError):
                errors.append("Line evidence cannot be read")
        elif kind == "image_region":
            from PIL import Image
            try:
                with Image.open(contained(ROOT, item["path"])) as img:
                    box = item.get("box")
                    if not isinstance(box, list) or len(box) != 4 or any(type(v) is not int for v in box) or not (0 <= box[0] < box[2] <= img.width and 0 <= box[1] < box[3] <= img.height):
                        errors.append("Pixel evidence box is outside actual image")
                    if not isinstance(item.get("field"), str) or not item["field"]:
                        errors.append("Pixel evidence needs field name")
            except (OSError, ValueError):
                errors.append("Pixel evidence image cannot be read")
        else:
            errors.append("Unsupported evidence kind")
    if manifest["module"] == "M04" and manifest["case_id"] == "normal":
        expected = read(contained(ROOT, manifest["expected"]))["facts"]["requirements"]
        required = {(r["line"], r["quote"]) for r in expected}
        seen = {(e.get("line"), e.get("quote")) for e in evidence if isinstance(e, dict) and e.get("kind") == "line"}
        if not required <= seen:
            errors.append("All 12 requirements need matching raw-line evidence")
    if manifest["module"] == "M08" and manifest["case_id"] == "normal":
        request = read(contained(ROOT, manifest["input"]))
        for asset in request["payload"]["assets"]:
            fields = {e.get("field") for e in evidence if isinstance(e, dict) and e.get("path") == asset["image"] and e.get("kind") == "image_region"}
            if not {"person_id", "certificate_id", "qualification", "valid_until"} <= fields:
                errors.append("Each clear certificate needs four actual pixel-region evidence records")
    return errors


def image_sequence(document):
    """Visible image relationships in document order, not merely ZIP media set."""
    hashes = []
    for node in document._element.iter(qn("a:blip")):
        rid = node.get(qn("r:embed"))
        if rid not in document.part.rels:
            raise ValueError("Visible image points to a missing relationship")
        relation = document.part.rels[rid]
        if relation.is_external or not relation.reltype.endswith("/image"):
            raise ValueError("Visible image must refer to an embedded image")
        hashes.append(hashlib.sha256(relation.target_part.blob).hexdigest())
    return hashes


def declared_artifact_paths(value, keys=()):
    """Walk path declarations, including nested groups and lists."""
    if isinstance(value, dict):
        for key, child in value.items():
            yield from declared_artifact_paths(child, keys + (key,))
    elif isinstance(value, list):
        for child in value:
            yield from declared_artifact_paths(child, keys)
    else:
        yield keys, value


def check_artifacts(manifest, result, result_path):
    spec = manifest["artifacts"]
    if not spec:
        return []
    directory = Path(result_path).parent.resolve()
    artifacts = result.get("artifacts", {})
    if not isinstance(artifacts, dict):
        return ["artifacts must be an object"]
    kind = spec["kind"]
    errors = []
    if kind in ("no_docx", "no_zip"):
        suffix, role, message = ((".docx", "document", "A blocked/unconfirmed case produced DOCX")
                                 if kind == "no_docx" else
                                 (".zip", "archive", "Nonpublic case produced a public ZIP"))
        # Inspect actual filenames independently of OS glob case sensitivity.
        prohibited = any(p.is_file() and p.suffix.casefold() == suffix for p in directory.rglob("*"))
        for keys, relative in declared_artifact_paths(artifacts):
            try:
                path = contained(directory, relative)
            except ValueError as exc:
                errors.append("Invalid artifact declaration: " + str(exc))
                continue
            # A declared document/archive is forbidden even with a disguised
            # extension or before it is written. Also catch aliases by suffix.
            if role in keys or Path(relative).suffix.casefold() == suffix or path.suffix.casefold() == suffix:
                prohibited = True
        if prohibited:
            errors.append(message)
        return errors
    if kind == "personnel_docx":
        path = contained(directory, artifacts.get("document"))
        doc = Document(path)
        if len(doc.tables) != 1:
            return ["Expected exactly one personnel table"]
        rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
        errors += compare([spec["columns"]] + spec["rows"], rows, "actual_docx.table")
        actual_images = image_sequence(doc)
        expected_images = [digest(contained(ROOT, p)) for p in spec["image_paths"]]
        errors += compare(expected_images, actual_images, "actual_docx.visible_images")
        with zipfile.ZipFile(path) as z:
            media = {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith("word/media/")}
        if media != set(expected_images):
            errors.append("DOCX contains missing/old/unselected image bytes")
        text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
        if any(word in text for word in spec["excluded_text"]):
            errors.append("Old/unselected person remains in actual DOCX")
        if any(word not in text for word in spec["project_fields"]) or "{{" in text or "[TODO" in text:
            errors.append("Actual DOCX has wrong context or unresolved placeholders")
    elif kind == "public_zip":
        path = contained(directory, artifacts.get("archive"))
        allowed = {Path(p).name: p for p in spec["allowed_inputs"]}
        with zipfile.ZipFile(path) as z:
            if set(z.namelist()) != set(allowed):
                errors.append("Public ZIP must contain exactly the three approved-for-public candidate source files")
            for name in z.namelist():
                if name not in allowed or hashlib.sha256(z.read(name)).hexdigest() != digest(contained(ROOT, allowed[name])):
                    errors.append("Public ZIP content is not the registered allowed candidate: " + name)
    elif kind == "rerun_history":
        request = read(contained(ROOT, manifest["input"]))["payload"]
        old = contained(directory, artifacts.get("old_run"))
        new = contained(directory, artifacts.get("new_run"))
        if old == new or old.is_relative_to(new) or new.is_relative_to(old):
            errors.append("Old and new run directories must be disjoint")
        for module, item in request["previous_outputs"].items():
            if digest(old / f"{module}.json") != item["sha256"]:
                errors.append("Historical output was overwritten: " + module)
        for module in result["facts"]["unchanged"]:
            if digest(new / f"{module}.json") != request["previous_outputs"][module]["sha256"]:
                errors.append("Unchanged module was rewritten: " + module)
        for module in result["facts"]["rerun"]:
            output = read(new / f"{module}.json")
            if output.get("run_id") != request["new_run_id"]:
                errors.append("Rerun result retains old run ID: " + module)
        if list(new.glob("*release*.json")):
            for p in new.glob("*release*.json"):
                if read(p).get("status") == "RELEASED":
                    errors.append("New changed run inherited an old release")
    elif kind == "rerun_scenarios":
        for name, facts in result["facts"]["scenarios"].items():
            nested = dict(manifest, artifacts={"kind": "rerun_history"})
            nested_result = dict(result, facts=facts, artifacts=artifacts.get("scenarios", {}).get(name, {}))
            errors += [name + ": " + e for e in check_artifacts(nested, nested_result, result_path)]
    else:
        errors.append("Unknown artifact inspection kind: " + kind)
    return errors


def check_result(manifest, result, result_path):
    errors = verify_case(manifest)
    expected = read(contained(ROOT, manifest["expected"]))
    if not isinstance(result, dict):
        return errors + ["result.json must be an object"]
    for key, value in (("schema_version", RESULT_VERSION), ("module", manifest["module"]), ("case_id", manifest["case_id"]), ("source_mode", "DEMO"), ("status", expected["status"]), ("real_submission_authorized", False)):
        errors += compare(value, result.get(key), key)
    errors += compare(expected["facts"], result.get("facts"), "facts")
    codes = result.get("issue_codes")
    if not isinstance(codes, list) or any(not isinstance(v, str) for v in codes):
        errors.append("issue_codes must be an array of strings")
    elif not set(expected["issue_codes"]) <= set(codes) or (not expected["issue_codes"] and codes):
        errors.append(f"issue_codes: expected {expected['issue_codes']}, received {codes}")
    hashes = {i["path"]: i["sha256"] for i in manifest["inputs"]}
    errors += compare(hashes, result.get("input_hashes"), "input_hashes")
    try:
        errors += check_evidence(manifest, result)
        errors += check_artifacts(manifest, result, result_path)
    except (OSError, ValueError, KeyError, TypeError, PackageNotFoundError, zipfile.BadZipFile) as exc:
        errors.append("Actual artifact/evidence inspection failed: " + str(exc))
    return errors


def run_suite(module, adapter, output):
    failures, _ = verify_kit(module)
    if failures:
        raise ValueError("Corpus verification failed: " + "; ".join(failures))
    adapter = Path(adapter).resolve()
    if not adapter.is_file():
        raise ValueError("Adapter is not implemented or not found: " + str(adapter))
    output = Path(output).resolve()
    if output.exists():
        raise ValueError("Output already exists; preserve previous evidence and use a new directory")
    output.mkdir(parents=True)
    reports = []
    for manifest in cases(module):
        target = output / manifest["case_id"]; target.mkdir()
        command = [sys.executable, str(adapter), "--input", str(contained(ROOT, manifest["input"])), "--output", str(target)]
        try:
            process = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, timeout=180)
            (target / "stdout.txt").write_text(process.stdout, encoding="utf-8")
            (target / "stderr.txt").write_text(process.stderr, encoding="utf-8")
            result = target / "result.json"
            errors = check_result(manifest, read(result), result) if result.is_file() else ["Adapter did not produce result.json"]
            if process.returncode not in (0, 2):
                errors.append(f"Unexpected process exit code {process.returncode}")
            reports.append({"case_id": manifest["case_id"], "command": command, "exit_code": process.returncode, "passed": not errors, "errors": errors})
        except (OSError, ValueError, subprocess.TimeoutExpired) as exc:
            reports.append({"case_id": manifest["case_id"], "command": command, "passed": False, "errors": [str(exc)]})
    report = {"module": module, "kind": "STUDENT_IMPLEMENTATION_TEST", "passed": all(r["passed"] for r in reports), "cases": reports, "independent_review": "PENDING", "business_uat": "PENDING"}
    (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return report["passed"]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="mode", required=True)
    verify = sub.add_parser("verify-kit"); verify.add_argument("--module")
    check = sub.add_parser("check"); check.add_argument("--module", required=True); check.add_argument("--case", required=True); check.add_argument("--result", required=True)
    run = sub.add_parser("run"); run.add_argument("--module", required=True); run.add_argument("--adapter", required=True); run.add_argument("--output", required=True)
    args = parser.parse_args()
    try:
        if args.mode == "verify-kit":
            errors, count = verify_kit(args.module)
            print(json.dumps({"kind": "CORPUS_INTEGRITY_ONLY", "cases": count, "passed": not errors, "errors": errors, "student_implementation": "NOT_TESTED"}, ensure_ascii=False, indent=2))
            return 1 if errors else 0
        if args.mode == "check":
            manifest = next((c for c in cases(args.module) if c["case_id"] == args.case), None)
            if not manifest:
                raise ValueError("Unknown case ID")
            errors = check_result(manifest, read(args.result), args.result)
            print(json.dumps({"case": args.case, "passed": not errors, "errors": errors, "independent_review": "PENDING"}, ensure_ascii=False, indent=2))
            return 1 if errors else 0
        return 0 if run_suite(args.module, args.adapter, args.output) else 1
    except (OSError, ValueError, KeyError) as exc:
        print("Acceptance check failed: " + str(exc), file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
